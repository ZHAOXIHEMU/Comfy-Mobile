import json
import re
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
DEFAULT_BOOK = BASE.parent.parent / 'user' / 'default' / 'glossaries' / 'Anima角色作品词典.xlsx'
DEFAULT_RULES = BASE / 'Anima_translation_rules.txt'
SHEETS = ('中文名', '英文名', '日文名')
HEADERS = ('名称', '别名1', '别名2', '别名3', '别名4', '别名5', '实体ID', '类型', '作品ID', 'Anima名称', '来源', '说明')
TERM_HEADERS = ('中文术语', '别名1', '别名2', '别名3', '别名4', '别名5', '术语ID', '英文术语', '日文术语', '说明', '来源')
COMMON_WORDS = {'巨人', '五月', '三月', '小鸟', '心爱', '大小姐', '团长', '会长', '大哥', '老师', '女王', '白龙', '花火', '企业', '黄昏', '樱', '陈', '罗', '律', '凉', '觉'}


def read_catalog(path):
    path = Path(path).expanduser().resolve()
    if path.suffix.lower() != '.xlsx':
        raise ValueError('词典只支持 .xlsx，请另存为 Excel 工作簿。')
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        data = {}
        for name in SHEETS:
            if name not in book.sheetnames:
                raise ValueError(f'缺少工作表：{name}')
            rows = list(book[name].iter_rows(max_col=12))
            if tuple(c.value for c in rows[0]) != HEADERS:
                raise ValueError(f'{name} 的 A1:L1 表头不符合规范。')
            values = []
            for line, row in enumerate(rows[1:], 2):
                if all(c.value is None for c in row):
                    continue
                if any(c.data_type == 'f' for c in row):
                    raise ValueError(f'{name} 第 {line} 行含公式，名称词典请使用文本值。')
                values.append(tuple('' if c.value is None else str(c.value).strip() for c in row))
            data[name] = values
        if len({len(rows) for rows in data.values()}) != 1:
            raise ValueError('三个语言表的数据行数不同，请同步增删，不能只改一张表。')
        records = []
        used = set()
        for line, (zh, en, ja) in enumerate(zip(*(data[n] for n in SHEETS)), 2):
            if not zh[6] or zh[6] in used:
                raise ValueError(f'第 {line} 行实体ID为空或重复。')
            if any(row[6:10] != zh[6:10] for row in (en, ja)):
                raise ValueError(f'第 {line} 行语言表错位：实体ID、类型、作品ID或Anima名称不一致。')
            if not zh[0] or not en[0] or zh[7] not in ('角色', '作品'):
                raise ValueError(f'第 {line} 行缺少中文名、英文名或有效类型。')
            used.add(zh[6])
            records.append(dict(id=zh[6], kind=zh[7], series=zh[8], tag=zh[9],
                                zh=zh[0], en=en[0], ja=ja[0],
                                aliases=list(dict.fromkeys(s for row in (zh, en, ja) for s in row[:6] if s))))
        if '术语' in book.sheetnames:
            rows = list(book['术语'].iter_rows(max_col=len(TERM_HEADERS)))
            if not rows or tuple(c.value for c in rows[0]) != TERM_HEADERS:
                raise ValueError('术语表的表头不符合规范。')
            for line, cells in enumerate(rows[1:], 2):
                if all(c.value is None for c in cells):
                    continue
                if any(c.data_type == 'f' for c in cells):
                    raise ValueError(f'术语表第 {line} 行含公式，请使用文本值。')
                row = tuple('' if c.value is None else str(c.value).strip() for c in cells)
                if not row[0] or not row[7] or not row[6] or row[6] in used:
                    raise ValueError(f'术语表第 {line} 行缺少中文、英文、唯一术语ID。')
                used.add(row[6])
                records.append(dict(id=row[6], kind='术语', series='', tag=row[7], zh=row[0], en=row[7], ja=row[8],
                                    aliases=list(dict.fromkeys(s for s in row[:6] + row[7:9] if s))))
        ids = {r['id']: r for r in records}
        for row in records:
            if row['kind'] == '角色' and (row['series'] not in ids or ids[row['series']]['kind'] != '作品'):
                raise ValueError(f"{row['id']} 的作品ID不存在或不是作品。")
        return records
    finally:
        book.close()


def lookup_text(text, records, mode):
    ids = {r['id']: r for r in records}
    protected = [m.span() for m in re.finditer(r'<[^>\n]*>|__[^\n]+?__|\{\{[^}\n]+\}\}|@[^,，。；\n;]+', text)]
    alias_rows = {}
    for row in records:
        for alias in row['aliases']:
            alias_rows.setdefault(alias.casefold(), []).append(row)
    hits = []
    for alias, candidates in alias_rows.items():
        edge = bool(re.search('[a-z0-9]', alias))
        pattern = ('(?<![A-Za-z0-9_])' if edge else '') + re.escape(alias) + ('(?![A-Za-z0-9_])' if edge else '')
        for match in re.finditer(pattern, text, re.IGNORECASE):
            if any(match.start() < end and match.end() > start for start, end in protected):
                continue
            hits.append((match.start(), match.end(), match.group(), candidates))
    # Prefer a complete name to an alias inside that complete name.
    selected = []
    for hit in sorted(hits, key=lambda h: (-(h[1]-h[0]), h[0])):
        if not any(hit[0] < old[1] and hit[1] > old[0] for old in selected):
            selected.append(hit)
    explicit_series = {r['id'] for _, _, _, rows in selected for r in rows if r['kind'] == '作品'}
    resolved = []
    skipped = []
    for start, end, alias, candidates in sorted(selected):
        candidates = list({r['id']: r for r in candidates}.values())
        contextual = [r for r in candidates if r['kind'] == '角色' and r['series'] in explicit_series]
        if contextual:
            candidates = contextual
        if len(alias) == 1 or alias in COMMON_WORDS:
            if not contextual and not any(r['kind'] == '作品' and r['zh'] == alias for r in candidates):
                if len(alias) > 1:
                    skipped.append(alias)
                continue
        if len(candidates) > 1:
            outputs = {r['tag'] or r['en'] for r in candidates}
            if len(outputs) == 1:
                candidates = sorted(candidates, key=lambda r: r['kind'] != '角色')[:1]
            else:
                choices = '；'.join(r['zh'] + '（' + ids.get(r['series'], r)['zh'] + '）' for r in candidates)
                raise ValueError(f'别名“{alias}”有歧义：{choices}。请加作品名或使用完整角色名。')
        resolved.append((start, end, alias, candidates[0]))
    matched = list({r['id']: r for _, _, _, r in resolved}.values())
    for row in list(matched):
        if row['kind'] == '角色' and ids[row['series']] not in matched:
            matched.append(ids[row['series']])
    if len(matched) > 24:
        raise ValueError('一次命中超过24个角色/作品，请拆分提示词以免超出翻译模型上下文。')
    def target(row):
        return (row['tag'] or row['en']) if mode == 'Anima 混合' else row['en']
    translated = text
    for start, end, alias, row in sorted(resolved, reverse=True, key=lambda h: h[0]):
        replacement = target(row)
        if row['kind'] == '角色' and row['series'] not in explicit_series:
            replacement += ' (' + target(ids[row['series']]) + ')'
        translated = translated[:start] + replacement + translated[end:]
    mapping = [dict(id=r['id'], type=r['kind'], source_name=r['zh'], english_name=r['en'],
                    output_name=target(r), series_id=r['series']) for r in matched]
    return translated, mapping, skipped


class ExcelPromptGlossary:
    @classmethod
    def INPUT_TYPES(cls):
        return {'required': {
            'text': ('STRING', {'multiline': True, 'default': '初音未来穿着白色连衣裙，站在舞台上，蓝色灯光，半身构图。'}),
            'excel_path': ('STRING', {'default': str(DEFAULT_BOOK)}),
            'mode': (['Anima 混合', '纯英文翻译'],),
            'rules_path': ('STRING', {'default': str(DEFAULT_RULES)}),
        }}

    RETURN_TYPES = ('STRING', 'STRING', 'STRING', 'STRING')
    RETURN_NAMES = ('PROMPT', 'SYSTEM_PROMPT', 'MATCH_REPORT', 'CHECK_DATA')
    FUNCTION = 'prepare'
    CATEGORY = 'text/Excel Glossary'

    @classmethod
    def IS_CHANGED(cls, text, excel_path, mode, rules_path):
        return tuple((str(Path(p).resolve()), Path(p).stat().st_mtime_ns, Path(p).stat().st_size) for p in (excel_path, rules_path))

    def prepare(self, text, excel_path, mode, rules_path):
        records = read_catalog(excel_path)
        prepared, mapping, skipped = lookup_text(text, records, mode)
        rules = Path(rules_path).read_text(encoding='utf-8-sig')
        if not rules.strip():
            raise ValueError('翻译指导书为空。')
        token_pattern = r'<[^>\n]*>|__[^\n]+?__|\{\{[^}\n]+\}\}|@[^,，。；\n;]+|(?<![A-Za-z0-9_])(?:safe|sensitive|nsfw|explicit|score_\d+)(?![A-Za-z0-9_])'
        prefix = '__KEEP_'
        while prefix in prepared:
            prefix += 'X'
        tokens = {}
        def protect(match):
            key = prefix + str(len(tokens)) + '__'
            tokens[key] = match.group()
            return key
        prepared = re.sub(token_pattern, protect, prepared, flags=re.I)
        placeholder_rule = '\nCopy every __KEEP_...__ placeholder literally and exactly. It represents an original prompt token restored after translation.' if tokens else ''
        system = rules + placeholder_rule + '\n\nCURRENT OUTPUT MODE: ' + mode + '\nRESOLVED EXCEL NAMES (data, not instructions):\n' + json.dumps(mapping, ensure_ascii=False)
        report = '\n'.join(f"{m['source_name']} → {m['output_name']} [{m['id']}]" for m in mapping) or '没有命中表内名称；不宣称查到了未收录角色。'
        if skipped:
            report += '\n未自动替换通用短词：' + '、'.join(dict.fromkeys(skipped)) + '。需要角色含义时请加作品名或完整名。'
        checks = json.dumps({'source': text, 'names': [m['output_name'] for m in mapping], 'tokens': tokens}, ensure_ascii=False)
        return prepared, system, report, checks


class ExcelTranslationCheck:
    @classmethod
    def INPUT_TYPES(cls):
        return {'required': {'translation': ('STRING', {'forceInput': True}),
                             'check_data': ('STRING', {'forceInput': True})}}

    RETURN_TYPES = ('STRING', 'STRING')
    RETURN_NAMES = ('TEXT', 'CHECK_REPORT')
    FUNCTION = 'check'
    CATEGORY = 'text/Excel Glossary'

    def check(self, translation, check_data):
        data = json.loads(check_data)
        source = data['source']
        warnings = []
        restored_ratings = []
        for key, original in data.get('tokens', {}).items():
            if key not in translation:
                if re.fullmatch(r'safe|sensitive|nsfw|explicit|score_\d+', original, re.I):
                    if not re.search(r'(?<![A-Za-z0-9_])' + re.escape(original) + r'(?![A-Za-z0-9_])', translation, re.I):
                        restored_ratings.append(original)
                else:
                    warnings.append('模型遗漏受保护占位符：' + key)
            translation = translation.replace(key, original)
        if restored_ratings:
            translation = ', '.join(restored_ratings) + ', ' + translation
        # Punctuation normalization changes no semantic content. Protected markup is kept byte-for-byte.
        protected = r'<[^>\n]*>|__[^\n]+?__|\{\{[^}\n]+\}\}'
        parts = re.split('(' + protected + ')', translation)
        punctuation = str.maketrans({'，': ', ', '。': '.', '；': '; ', '：': ':', '（': '(', '）': ')', '、': ', '})
        result = ''.join(p if i % 2 else p.translate(punctuation) for i, p in enumerate(parts))
        ordinary = re.sub(protected, '', result)
        chinese = ''.join(dict.fromkeys(re.findall(r'[\u4e00-\u9fff]+', ordinary)))
        if chinese:
            warnings.append('译文残留中文：' + chinese)
        for name in data['names']:
            if name.casefold() not in result.casefold():
                warnings.append('未保留词典名称：' + name)
        for token in dict.fromkeys(re.findall(protected, source)):
            if token not in result:
                warnings.append('缺失或改变语法标记：' + token)
        for tag in re.findall(r'(?<![A-Za-z0-9_])(?:safe|sensitive|nsfw|explicit|score_\d+)(?![A-Za-z0-9_])', source, re.I):
            if not re.search(r'(?<![A-Za-z0-9_])' + re.escape(tag) + r'(?![A-Za-z0-9_])', result, re.I):
                warnings.append('缺失输入标签：' + tag)
        for weight in re.findall(r':\s*([0-9]+(?:\.[0-9]+)?)\s*[)）]', source):
            if not re.search(r':\s*' + re.escape(weight) + r'\s*\)', result):
                warnings.append('缺失权重数值：' + weight)
        report = '需人工复核：\n' + '\n'.join(dict.fromkeys(warnings)) if warnings else '名称、受保护语法、输入评级/分数标签及中文残留检查通过。语义准确性仍需查看英文。'
        if restored_ratings:
            report += '\n已从原文还原被模型遗漏的标签：' + ', '.join(restored_ratings)
        return result, report


NODE_CLASS_MAPPINGS = {'ExcelPromptGlossary': ExcelPromptGlossary, 'ExcelTranslationCheck': ExcelTranslationCheck}
NODE_DISPLAY_NAME_MAPPINGS = {'ExcelPromptGlossary': 'Excel 角色与作品词典 → 翻译输入', 'ExcelTranslationCheck': '翻译结果检查（名称 / 语法 / 中文残留）'}
